#!/usr/bin/env python3
"""Regression tests for SellerSprite keyword and candidate discovery."""

from __future__ import annotations

import csv
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl


def write_keyword_export(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append([f"Column {index}" for index in range(36)])
    for keyword, volume, asin in (
        ("ai voice recorder with transcription", 4500, "B000000001"),
        ("voice recorder case", 1200, "B000000099"),
    ):
        row: list[object] = [None] * 36
        row[0] = keyword
        row[6] = volume
        row[9] = 100
        row[10] = 0.02
        row[15] = 500
        row[20] = "$2.50"
        row[26] = asin
        sheet.append(row)
    workbook.save(path)


def product_row(child: str, parent: str, brand: str, title: str, category: str) -> list[object]:
    row: list[object] = [None] * 65
    row[0] = child
    row[1] = "Color: Black"
    row[2] = "Memory Storage Capacity:64 GB | Compatible Devices:Smartphone | Format:WAV"
    row[3] = brand
    row[5] = title
    row[6] = f"https://www.amazon.com/dp/{child}"
    row[7] = f"https://images.example/{child}.jpg"
    row[8] = parent
    row[9] = category
    row[14] = category.split(":")[-1]
    row[16] = 100
    row[19] = 15000
    row[23] = 150
    row[27] = 50
    row[29] = 4.5
    row[34] = "2026-01-01"
    row[36] = "FBA"
    return row


def write_competitor_export(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append([f"Column {index}" for index in range(65)])
    sheet.append(
        product_row(
            "B000000001",
            "P000000001",
            "Direct",
            "AI Voice Recorder with Transcription and Summary",
            "Electronics:Digital Voice Recorders",
        )
    )
    sheet.append(
        product_row(
            "B000000002",
            "P000000002",
            "Traditional",
            "Voice Recorder with AI Noise Reduction",
            "Electronics:Digital Voice Recorders",
        )
    )
    sheet.append(
        product_row(
            "B000000003",
            "P000000003",
            "Laptop",
            "Laptop with AI Voice Recorder bundle",
            "Electronics:Laptops",
        )
    )
    workbook.save(path)


def write_detail(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["ASIN", "Parent", "", "", "", "Title", "", "2026-08"])
    sheet.append(["B000000001", "P000000001", "", "", "", "Direct", "", 100])
    sheet.append(["B000000002", "P000000002", "", "", "", "Adjacent", "", 50])
    workbook.save(path)


class DiscoveryTest(unittest.TestCase):
    def test_keyword_clusters_and_candidate_layers(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            keyword = root / "keyword.xlsx"
            competitor = root / "competitor.xlsx"
            detail = root / "detail.xlsx"
            rules = root / "rules.json"
            output = root / "output"
            write_keyword_export(keyword)
            write_competitor_export(competitor)
            write_detail(detail)
            rules.write_text(
                json.dumps(
                    {
                        "required_category_terms": ["digital voice recorders"],
                        "exclude_category_terms": ["laptops"],
                        "core_title_terms": ["transcription", "summary"],
                        "adjacent_title_terms": ["noise reduction"],
                        "exclude_title_terms": ["bundle"],
                    }
                ),
                encoding="utf-8",
            )
            script = Path(__file__).with_name("discover_sellersprite_market.py")
            result = subprocess.run(
                [
                    sys.executable,
                    str(script),
                    "--keyword-export",
                    str(keyword),
                    "--competitor-export",
                    str(competitor),
                    "--detail",
                    str(detail),
                    "--rules",
                    str(rules),
                    "--seed-keyword",
                    "ai voice recorder",
                    "--seed-asin",
                    "B000000001",
                    "--output-dir",
                    str(output),
                ],
                check=True,
                capture_output=True,
                text=True,
            )
            self.assertIn("targets=1 children / 1 parents", result.stdout)
            with (output / "candidate_manifest.csv").open(encoding="utf-8-sig", newline="") as handle:
                rows = {row["child_asin"]: row for row in csv.DictReader(handle)}
            self.assertEqual(rows["B000000001"]["decision"], "core_direct")
            self.assertEqual(rows["B000000002"]["decision"], "adjacent_substitute")
            self.assertEqual(rows["B000000003"]["decision"], "excluded_noise")
            self.assertEqual(rows["B000000001"]["memory_storage"], "64 GB")
            with (output / "keyword_evidence.csv").open(encoding="utf-8-sig", newline="") as handle:
                keywords = {row["keyword"]: row for row in csv.DictReader(handle)}
            self.assertEqual(keywords["ai voice recorder with transcription"]["cluster"], "AI outcome")
            self.assertEqual(keywords["voice recorder case"]["relevant"], "no")
            summary = json.loads((output / "discovery_summary.json").read_text(encoding="utf-8"))
            self.assertEqual(summary["target_parent_count"], 1)


if __name__ == "__main__":
    unittest.main()
