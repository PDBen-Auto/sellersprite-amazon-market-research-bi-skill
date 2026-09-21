#!/usr/bin/env python3
"""Build an auditable keyword and ASIN discovery manifest from SellerSprite exports."""

from __future__ import annotations

import argparse
import csv
import json
import math
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Iterable

import openpyxl


ASIN_RE = re.compile(r"^[A-Z0-9]{10}$", re.IGNORECASE)
MONTH_RE = re.compile(r"^\d{4}-\d{2}(?:\(\$\))?$")

DEFAULT_RULES: dict[str, Any] = {
    "required_category_terms": [],
    "core_title_terms": [],
    "adjacent_title_terms": [],
    "exclude_title_terms": ["case", "cover", "charger", "adapter", "replacement"],
    "exclude_category_terms": [],
    "core_brands": [],
    "include_asins": [],
    "exclude_asins": [],
    "overrides": {},
}

KEYWORD_FIELDS = [
    "keyword",
    "cluster",
    "relevant",
    "relevance_reason",
    "search_volume",
    "search_growth_3m",
    "search_growth_yoy",
    "purchases",
    "purchase_rate",
    "product_count",
    "ppc_bid",
    "bid_range",
    "top_asins",
]

CANDIDATE_FIELDS = [
    "child_asin",
    "parent_asin",
    "decision",
    "relevance_score",
    "decision_reason",
    "detail_available",
    "brand",
    "title",
    "category",
    "leaf_category",
    "sku",
    "color",
    "memory_storage",
    "compatible_devices",
    "format",
    "product_weight",
    "product_dimensions",
    "package_weight",
    "current_units",
    "current_sales_amount",
    "price",
    "rating",
    "reviews",
    "launch_date",
    "fulfillment",
    "url",
    "image",
    "attributes_raw",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--keyword-export", required=True)
    parser.add_argument("--competitor-export", required=True)
    parser.add_argument("--detail", help="Optional SellerSprite detailed export used as an availability gate")
    parser.add_argument("--rules", help="JSON file containing project-specific inclusion and exclusion rules")
    parser.add_argument("--seed-keyword", default="")
    parser.add_argument("--seed-asin", default="")
    parser.add_argument("--output-dir", required=True)
    parser.add_argument(
        "--allow-without-detail",
        action="store_true",
        help="Put core candidates without detailed-export coverage in target ASIN files",
    )
    return parser.parse_args()


def cell(row: tuple[Any, ...], index: int) -> Any:
    return row[index] if index < len(row) else None


def text(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return None
        return float(value)
    cleaned = str(value).strip().replace(",", "").replace("$", "")
    if not cleaned:
        return None
    try:
        return float(cleaned)
    except ValueError:
        return None


def asin(value: Any) -> str:
    candidate = text(value).upper()
    return candidate if ASIN_RE.fullmatch(candidate) else ""


def terms(values: Iterable[Any]) -> list[str]:
    return [text(value).casefold() for value in values if text(value)]


def matches_any(haystack: str, needles: Iterable[str]) -> list[str]:
    folded = haystack.casefold()
    return [needle for needle in needles if needle and needle in folded]


def load_rules(path: str | None) -> dict[str, Any]:
    rules = dict(DEFAULT_RULES)
    if path:
        supplied = json.loads(Path(path).read_text(encoding="utf-8"))
        if not isinstance(supplied, dict):
            raise ValueError("Rules JSON must contain an object")
        rules.update(supplied)
    for key in (
        "required_category_terms",
        "core_title_terms",
        "adjacent_title_terms",
        "exclude_title_terms",
        "exclude_category_terms",
        "core_brands",
        "include_asins",
        "exclude_asins",
    ):
        rules[key] = terms(rules.get(key, []))
    rules["overrides"] = {
        str(key).upper(): value for key, value in (rules.get("overrides") or {}).items()
    }
    return rules


def write_csv(path: Path, fieldnames: list[str], rows: Iterable[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def split_top_asins(row: tuple[Any, ...]) -> list[str]:
    values = [cell(row, 26), cell(row, 29), cell(row, 32)]
    values.extend(text(cell(row, 35)).split(","))
    result: list[str] = []
    for value in values:
        candidate = asin(value)
        if candidate and candidate not in result:
            result.append(candidate)
    return result


def keyword_cluster(keyword: str) -> tuple[str, bool, str]:
    folded = keyword.casefold()
    if matches_any(folded, ("case", "cover", "charger", "adapter", "replacement", "manual")):
        return "accessory/noise", False, "Accessory or support intent"
    if matches_any(folded, ("plaud", "notepin", "soundcore", "hidock", "pocket ai", "comulytic")):
        return "brand/product", True, "Named brand or product family"
    if matches_any(folded, ("no subscription", "free transcription", "unlimited", "offline")):
        return "commercial constraint", True, "Plan, subscription, or offline constraint"
    if matches_any(folded, ("meeting", "lecture", "college", "interview", "classroom", "call")):
        return "use case", True, "Recording use-case intent"
    if matches_any(folded, ("transcri", "summar", "note tak", "notetaker", "transcription device")):
        return "AI outcome", True, "Transcription, summary, or note-taking outcome"
    if matches_any(folded, ("wearable", "magnetic", "pocket", "pen", "card")):
        return "form factor", True, "Physical form-factor intent"
    if matches_any(folded, ("voice recorder", "ai recorder", "digital recorder")):
        return "generic category", True, "Generic recorder intent"
    return "review", False, "Does not clearly express the target recording job"


def read_keywords(path: Path) -> list[dict[str, Any]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    rows: list[dict[str, Any]] = []
    for raw in sheet.iter_rows(min_row=2, values_only=True):
        keyword = text(cell(raw, 0))
        if not keyword:
            continue
        cluster, relevant, reason = keyword_cluster(keyword)
        rows.append(
            {
                "keyword": keyword,
                "cluster": cluster,
                "relevant": "yes" if relevant else "no",
                "relevance_reason": reason,
                "search_volume": number(cell(raw, 6)),
                "search_growth_3m": number(cell(raw, 7)),
                "search_growth_yoy": number(cell(raw, 8)),
                "purchases": number(cell(raw, 9)),
                "purchase_rate": number(cell(raw, 10)),
                "product_count": number(cell(raw, 15)),
                "ppc_bid": number(cell(raw, 20)),
                "bid_range": text(cell(raw, 21)),
                "top_asins": " ".join(split_top_asins(raw)),
            }
        )
    rows.sort(key=lambda row: (-(row["search_volume"] or 0), row["keyword"]))
    return rows


def parse_attributes(raw: str) -> dict[str, str]:
    parsed: dict[str, str] = {}
    for part in raw.split("|"):
        key, separator, value = part.partition(":")
        if separator and key.strip() and value.strip():
            parsed[key.strip().casefold()] = value.strip()
    return parsed


def detail_asins(path: Path | None) -> set[str]:
    if path is None:
        return set()
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    result: set[str] = set()
    for sheet in workbook.worksheets:
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        if not any(MONTH_RE.fullmatch(text(value)) for value in header):
            continue
        for raw in sheet.iter_rows(min_row=2, values_only=True):
            candidate = next((asin(value) for value in raw[:7] if asin(value)), "")
            if candidate:
                result.add(candidate)
        if result:
            break
    return result


def classify_candidate(row: dict[str, Any], rules: dict[str, Any], seed_asin: str) -> tuple[str, int, str]:
    child = row["child_asin"]
    parent = row["parent_asin"]
    title_value = row["title"].casefold()
    category_value = row["category"].casefold()
    brand_value = row["brand"].casefold()
    override = rules["overrides"].get(child) or rules["overrides"].get(parent)
    if override:
        if isinstance(override, str):
            return override, 100, "Project override"
        return (
            override.get("decision", "review"),
            int(override.get("score", 100)),
            override.get("reason", "Project override"),
        )

    if child.casefold() in rules["exclude_asins"] or parent.casefold() in rules["exclude_asins"]:
        return "excluded_noise", 0, "Explicit ASIN exclusion"
    if child.casefold() in rules["include_asins"] or parent.casefold() in rules["include_asins"]:
        return "core_direct", 100, "Explicit ASIN inclusion"
    if seed_asin and child == seed_asin:
        return "core_direct", 100, "Seed competitor ASIN"

    exclude_category = matches_any(category_value, rules["exclude_category_terms"])
    exclude_title = matches_any(title_value, rules["exclude_title_terms"])
    required_category = matches_any(category_value, rules["required_category_terms"])
    core_title = matches_any(title_value, rules["core_title_terms"])
    adjacent_title = matches_any(title_value, rules["adjacent_title_terms"])
    trusted_brand = brand_value in rules["core_brands"]

    if exclude_category:
        return "excluded_noise", 0, f"Excluded category: {', '.join(exclude_category)}"
    if exclude_title:
        return "excluded_noise", 5, f"Excluded title term: {', '.join(exclude_title)}"
    if rules["required_category_terms"] and not required_category:
        return "excluded_noise", 10, "Outside required category path"
    if core_title:
        score = 85 + min(10, 2 * len(core_title)) + (5 if trusted_brand else 0)
        return "core_direct", min(score, 100), f"Core outcome terms: {', '.join(core_title[:4])}"
    if trusted_brand:
        return "core_direct", 80, "Trusted direct-competitor brand within required category"
    if adjacent_title:
        return "adjacent_substitute", 45, f"Adjacent recorder terms: {', '.join(adjacent_title[:4])}"
    return "review", 30, "Category match without enough evidence of transcription or AI-note outcome"


def read_candidates(path: Path, rules: dict[str, Any], available: set[str], seed_asin: str) -> list[dict[str, Any]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = next((item for item in workbook.worksheets if item.max_column >= 30), workbook.worksheets[0])
    rows: list[dict[str, Any]] = []
    for raw in sheet.iter_rows(min_row=2, values_only=True):
        child = asin(cell(raw, 0))
        if not child:
            continue
        parent = asin(cell(raw, 8)) or child
        attributes_raw = text(cell(raw, 2))
        attributes = parse_attributes(attributes_raw)
        sku = text(cell(raw, 1))
        color = ""
        if sku.casefold().startswith("color:"):
            color = sku.split(":", 1)[1].strip()
        candidate: dict[str, Any] = {
            "child_asin": child,
            "parent_asin": parent,
            "detail_available": "yes" if not available or child in available else "no",
            "brand": text(cell(raw, 3)) or "Unknown",
            "title": text(cell(raw, 5)),
            "category": text(cell(raw, 9)),
            "leaf_category": text(cell(raw, 14)),
            "sku": sku,
            "color": color,
            "memory_storage": attributes.get("memory storage capacity", ""),
            "compatible_devices": attributes.get("compatible devices", ""),
            "format": attributes.get("format", ""),
            "product_weight": text(cell(raw, 56)) or text(cell(raw, 55)),
            "product_dimensions": text(cell(raw, 58)) or text(cell(raw, 57)),
            "package_weight": text(cell(raw, 60)) or text(cell(raw, 59)),
            "current_units": number(cell(raw, 16)),
            "current_sales_amount": number(cell(raw, 19)),
            "price": number(cell(raw, 23)),
            "rating": number(cell(raw, 29)),
            "reviews": number(cell(raw, 27)),
            "launch_date": text(cell(raw, 34)),
            "fulfillment": text(cell(raw, 36)),
            "url": text(cell(raw, 6)) or f"https://www.amazon.com/dp/{child}",
            "image": text(cell(raw, 7)),
            "attributes_raw": attributes_raw,
        }
        decision, score, reason = classify_candidate(candidate, rules, seed_asin)
        candidate.update(
            {"decision": decision, "relevance_score": score, "decision_reason": reason}
        )
        rows.append(candidate)
    rows.sort(
        key=lambda row: (
            {"core_direct": 0, "adjacent_substitute": 1, "review": 2, "excluded_noise": 3}.get(
                row["decision"], 4
            ),
            -(row["current_sales_amount"] or 0),
            row["child_asin"],
        )
    )
    return rows


def parent_manifest(candidates: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in candidates:
        grouped[row["parent_asin"]].append(row)
    result: list[dict[str, Any]] = []
    decision_priority = {"core_direct": 0, "adjacent_substitute": 1, "review": 2, "excluded_noise": 3}
    for parent, children in grouped.items():
        representative = max(
            children,
            key=lambda row: (row["current_sales_amount"] or 0, row["current_units"] or 0),
        )
        best = min(children, key=lambda row: decision_priority.get(row["decision"], 4))
        result.append(
            {
                "parent_asin": parent,
                "decision": best["decision"],
                "relevance_score": max(row["relevance_score"] for row in children),
                "decision_reason": best["decision_reason"],
                "detail_available": "yes" if any(row["detail_available"] == "yes" for row in children) else "no",
                "child_count": len({row["child_asin"] for row in children}),
                "child_asins": " ".join(sorted({row["child_asin"] for row in children})),
                "brand": representative["brand"],
                "title": representative["title"],
                "category": representative["category"],
                "current_units": representative["current_units"],
                "current_sales_amount": representative["current_sales_amount"],
                "price": representative["price"],
                "rating": representative["rating"],
                "reviews": representative["reviews"],
                "url": representative["url"],
                "image": representative["image"],
            }
        )
    result.sort(
        key=lambda row: (
            decision_priority.get(row["decision"], 4),
            -(row["current_sales_amount"] or 0),
            row["parent_asin"],
        )
    )
    return result


def main() -> None:
    args = parse_args()
    output_dir = Path(args.output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    rules = load_rules(args.rules)
    available = detail_asins(Path(args.detail).resolve() if args.detail else None)
    keywords = read_keywords(Path(args.keyword_export).resolve())
    candidates = read_candidates(
        Path(args.competitor_export).resolve(), rules, available, asin(args.seed_asin)
    )
    parents = parent_manifest(candidates)

    write_csv(output_dir / "keyword_evidence.csv", KEYWORD_FIELDS, keywords)
    write_csv(output_dir / "candidate_manifest.csv", CANDIDATE_FIELDS, candidates)
    write_csv(
        output_dir / "parent_scope_manifest.csv",
        [
            "parent_asin",
            "decision",
            "relevance_score",
            "decision_reason",
            "detail_available",
            "child_count",
            "child_asins",
            "brand",
            "title",
            "category",
            "current_units",
            "current_sales_amount",
            "price",
            "rating",
            "reviews",
            "url",
            "image",
        ],
        parents,
    )

    target_children = [
        row["child_asin"]
        for row in candidates
        if row["decision"] == "core_direct"
        and (args.allow_without_detail or row["detail_available"] == "yes")
    ]
    target_parents = [
        row["parent_asin"]
        for row in parents
        if row["decision"] == "core_direct"
        and (args.allow_without_detail or row["detail_available"] == "yes")
    ]
    (output_dir / "target_child_asins.txt").write_text(
        "\n".join(target_children) + ("\n" if target_children else ""), encoding="utf-8"
    )
    (output_dir / "target_parent_asins.txt").write_text(
        "\n".join(target_parents) + ("\n" if target_parents else ""), encoding="utf-8"
    )

    keyword_clusters = Counter(row["cluster"] for row in keywords if row["relevant"] == "yes")
    decision_counts = Counter(row["decision"] for row in candidates)
    summary = {
        "seed_keyword": args.seed_keyword,
        "seed_asin": asin(args.seed_asin),
        "keyword_rows": len(keywords),
        "relevant_keyword_rows": sum(row["relevant"] == "yes" for row in keywords),
        "keyword_clusters": dict(sorted(keyword_clusters.items())),
        "candidate_child_rows": len(candidates),
        "candidate_parent_rows": len(parents),
        "detail_child_rows": len(available),
        "decision_counts": dict(sorted(decision_counts.items())),
        "target_child_count": len(target_children),
        "target_parent_count": len(target_parents),
        "files": {
            "keyword_export": Path(args.keyword_export).name,
            "competitor_export": Path(args.competitor_export).name,
            "detail_export": Path(args.detail).name if args.detail else None,
            "rules": Path(args.rules).name if args.rules else None,
        },
        "limitations": [
            "Keyword results are discovery evidence, not category TAM.",
            "Candidate decisions are rule-assisted and remain auditable in the manifest.",
            "Target ASIN files are limited to detailed-export coverage unless --allow-without-detail is used.",
        ],
    }
    (output_dir / "discovery_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(
        f"Clustered {len(keywords)} keywords; classified {len(candidates)} child ASINs / "
        f"{len(parents)} parents; targets={len(target_children)} children / {len(target_parents)} parents."
    )


if __name__ == "__main__":
    main()
