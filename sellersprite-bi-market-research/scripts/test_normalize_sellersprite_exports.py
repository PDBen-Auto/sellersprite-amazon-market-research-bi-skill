#!/usr/bin/env python3
"""Synthetic regression tests for the SellerSprite BI normalizer."""

from __future__ import annotations

import csv
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl


MONTHS = [f"2025-{month:02d}" for month in range(1, 7)]


def add_detail_sheet(
    workbook: openpyxl.Workbook,
    title: str,
    rows: list[tuple[str, str, str, list[float | None]]],
    dollar: bool = False,
) -> None:
    sheet = workbook.create_sheet(title)
    headers = ["ASIN", "Parent ASIN", "", "", "", "Title", "", ""]
    headers.extend(f"{month}($)" if dollar else month for month in MONTHS)
    sheet.append(headers)
    for child, parent, product_title, values in rows:
        row = [child, parent, "", "", "", product_title, "", ""]
        row.extend(values)
        sheet.append(row)


def write_detail(path: Path) -> None:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    add_detail_sheet(
        workbook,
        "Sales",
        [
            ("B000000001", "P000000001", "Parent A child 1", [100, 100, 100, 100, 100, 100]),
            ("B000000002", "P000000001", "Parent A child 2", [80, None, None, None, None, None]),
            ("B000000003", "P000000002", "Parent B", [50, 50, 50, 50, 50, 50]),
            ("B000000004", "B000000004", "Parent C incomplete", [10, 10, 10, 10, 10, 10]),
        ],
    )
    add_detail_sheet(
        workbook,
        "Sales Amount",
        [
            ("B000000001", "P000000001", "Parent A child 1", [10, 10, 10, 30, 30, 30]),
            ("B000000002", "P000000001", "Parent A child 2", [9, None, None, None, None, None]),
            ("B000000003", "P000000002", "Parent B", [20, 0, 20, 15, 15, 15]),
            ("B000000004", "B000000004", "Parent C incomplete", [100, 100, None, 100, 100, 100]),
        ],
        dollar=True,
    )
    add_detail_sheet(
        workbook,
        "Price",
        [
            ("B000000001", "P000000001", "Parent A child 1", [50, 50, 50, 50, 50, 50]),
            ("B000000002", "P000000001", "Parent A child 2", [55, 55, 55, 55, 55, 55]),
            ("B000000003", "P000000002", "Parent B", [40, 40, 40, 40, 40, 40]),
            ("B000000004", "B000000004", "Parent C incomplete", [60, 60, 60, 60, 60, 60]),
        ],
        dollar=True,
    )
    workbook.save(path)


def write_standard(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    sheet.append([f"Column {index}" for index in range(40)])
    rows = [
        ("B000000001", "P000000001", "Brand A", "Parent A child 1", 100, 5000),
        ("B000000002", "P000000001", "Brand A", "Parent A child 2", 80, 4400),
        ("B000000003", "P000000002", "Brand B", "Parent B", 50, 2000),
        ("B000000004", "P000000003", "Brand C", "Parent C incomplete", 10, 600),
    ]
    for child, parent, brand, title, units, revenue in rows:
        row: list[object] = [None] * 40
        row[0] = child
        row[3] = brand
        row[5] = title
        row[6] = f"https://www.amazon.com/dp/{child}"
        row[7] = f"https://images.example/{child}.jpg"
        row[8] = parent
        row[9] = "Automotive"
        row[16] = units
        row[19] = revenue
        row[23] = revenue / units
        row[27] = 100
        row[29] = 4.5
        row[34] = "2024-01-01"
        row[36] = "FBA"
        sheet.append(row)
    unrelated: list[object] = [None] * 40
    unrelated[0] = "B000000099"
    unrelated[3] = "Unrelated"
    unrelated[5] = "Standard-export row absent from the detailed selection"
    unrelated[8] = "P000000099"
    unrelated[16] = 999
    unrelated[19] = 99999
    sheet.append(unrelated)
    workbook.save(path)


class NormalizerTest(unittest.TestCase):
    def test_parent_dedupe_sales_amount_and_rank_momentum(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            detail = root / "detail.xlsx"
            standard = root / "standard.xlsx"
            output = root / "normalized"
            write_detail(detail)
            write_standard(standard)
            script = Path(__file__).with_name("normalize_sellersprite_exports.py")
            result = subprocess.run(
                [
                    sys.executable,
                    str(script),
                    "--detail",
                    str(detail),
                    "--standard",
                    str(standard),
                    "--output-dir",
                    str(output),
                    "--as-of",
                    "2025-07-15",
                    "--rank-window-months",
                    "3",
                ],
                check=True,
                capture_output=True,
                text=True,
            )
            self.assertIn("rank_denominator=2", result.stdout)

            expected = {
                "dashboard_data.json",
                "market_monthly.csv",
                "parent_monthly.csv",
                "parent_current.csv",
                "conflicts.csv",
                "rank_momentum.csv",
            }
            self.assertTrue(expected.issubset({path.name for path in output.iterdir()}))
            payload = json.loads((output / "dashboard_data.json").read_text(encoding="utf-8"))
            self.assertEqual(payload["metadata"]["unique_parents"], 3)
            self.assertEqual(payload["metadata"]["standard_child_rows"], 4)
            self.assertEqual(payload["metadata"]["parent_ids_reconciled_from_standard"], 1)
            self.assertEqual(payload["parent_current"][0]["fulfillment"], "FBA")
            self.assertTrue(all(row["launch_date"] == "2024-01-01" for row in payload["parent_current"]))

            january_a = next(
                row
                for row in payload["parent_monthly"]
                if row["parent_asin"] == "P000000001" and row["month"] == "2025-01"
            )
            self.assertEqual(january_a["units"], 100)
            self.assertEqual(january_a["seller_sprite_sales_amount"], 10)
            january_market = next(row for row in payload["market_monthly"] if row["month"] == "2025-01")
            self.assertEqual(january_market["units"], 160)

            february_b = next(
                row
                for row in payload["parent_monthly"]
                if row["parent_asin"] == "P000000002" and row["month"] == "2025-02"
            )
            self.assertEqual(february_b["seller_sprite_sales_amount"], 0)
            self.assertGreaterEqual(payload["data_quality"]["unit_conflict_count"], 1)
            self.assertGreaterEqual(payload["data_quality"]["sales_amount_conflict_count"], 1)

            summary = payload["rank_momentum"]["summary"]
            self.assertTrue(summary["available"])
            self.assertEqual(summary["global_denominator"], 2)
            ranks = {row["parent_asin"]: row for row in payload["rank_momentum"]["rows"]}
            self.assertNotIn("P000000003", ranks)
            self.assertEqual(ranks["P000000001"]["prior_rank"], 2)
            self.assertEqual(ranks["P000000001"]["recent_rank"], 1)
            self.assertEqual(ranks["P000000001"]["rank_change"], 1)
            self.assertEqual(ranks["P000000002"]["prior_rank"], 1)
            self.assertEqual(ranks["P000000002"]["recent_rank"], 2)
            self.assertEqual(ranks["P000000002"]["rank_change"], -1)

            with (output / "conflicts.csv").open(encoding="utf-8-sig", newline="") as handle:
                metrics = {row["metric"] for row in csv.DictReader(handle)}
            self.assertEqual(metrics, {"units", "seller_sprite_sales_amount"})

    def test_scope_manifest_filters_detail_and_preserves_scope_attributes(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            detail = root / "detail.xlsx"
            standard = root / "standard.xlsx"
            scope = root / "scope.csv"
            output = root / "normalized"
            write_detail(detail)
            write_standard(standard)
            with scope.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.DictWriter(
                    handle,
                    fieldnames=[
                        "child_asin",
                        "parent_asin",
                        "decision",
                        "relevance_score",
                        "decision_reason",
                    ],
                )
                writer.writeheader()
                writer.writerow(
                    {
                        "child_asin": "B000000001",
                        "parent_asin": "P000000001",
                        "decision": "core_direct",
                        "relevance_score": 95,
                        "decision_reason": "Transcription outcome",
                    }
                )
                writer.writerow(
                    {
                        "child_asin": "B000000003",
                        "parent_asin": "P000000002",
                        "decision": "adjacent_substitute",
                        "relevance_score": 40,
                        "decision_reason": "Traditional recorder",
                    }
                )
            script = Path(__file__).with_name("normalize_sellersprite_exports.py")
            subprocess.run(
                [
                    sys.executable,
                    str(script),
                    "--detail",
                    str(detail),
                    "--standard",
                    str(standard),
                    "--scope-manifest",
                    str(scope),
                    "--output-dir",
                    str(output),
                    "--as-of",
                    "2025-07-15",
                ],
                check=True,
                capture_output=True,
                text=True,
            )
            payload = json.loads((output / "dashboard_data.json").read_text(encoding="utf-8"))
            self.assertEqual(payload["metadata"]["unique_parents"], 1)
            self.assertEqual(payload["metadata"]["input_child_rows"], 2)
            row = payload["parent_current"][0]
            self.assertEqual(row["parent_asin"], "P000000001")
            self.assertEqual(row["market_layer"], "core_direct")
            self.assertEqual(row["relevance_score"], 95)

    def test_missing_sales_amount_sheet_is_explicitly_unavailable(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            detail = root / "detail-without-sales-amount.xlsx"
            output = root / "normalized"
            write_detail(detail)
            workbook = openpyxl.load_workbook(detail)
            workbook.remove(workbook["Sales Amount"])
            workbook.save(detail)
            script = Path(__file__).with_name("normalize_sellersprite_exports.py")
            subprocess.run(
                [
                    sys.executable,
                    str(script),
                    "--detail",
                    str(detail),
                    "--output-dir",
                    str(output),
                    "--as-of",
                    "2025-07-15",
                ],
                check=True,
                capture_output=True,
                text=True,
            )
            payload = json.loads((output / "dashboard_data.json").read_text(encoding="utf-8"))
            self.assertFalse(payload["rank_momentum"]["summary"]["available"])
            self.assertIsNone(payload["competition"]["brand_sales_amount"])
            self.assertTrue(
                any("No SellerSprite sales-amount sheet" in warning for warning in payload["data_quality"]["warnings"])
            )


if __name__ == "__main__":
    unittest.main()
